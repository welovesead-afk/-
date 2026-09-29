#!/usr/bin/env python3
"""제품명 통합표 만들기: 코드마다 공식명(카탈로그) 1개 + 다른 자료의 표기를 후보로 모음.
자동으로 합치지 않고 '후보(유사도)'로 표시 → 사람이 확정.

사용: python3 build_name_map.py <finished_goods_catalog.csv> <제품정보마스터.xlsx> <ledgers.csv> <출력 name_map.csv> <미대응 출력.csv>
"""
import collections
import csv
import difflib
import re
import sys

import openpyxl

SRC_ORDER = ['카탈로그', '제품분류표', '03 단품마스터', '04 선물세트마스터', '매출장부']


def norm(s):
    s = re.sub(r'\[[^\]]*\]|씨드|SEA\.?D|청정해역에서\s*자란|선물세트|세트', '', s, flags=re.I)
    s = s.replace('＆', '&').replace('조각', '자른').replace('건미역', '미역').replace('건다시마', '다시마').replace('G', 'g')
    return re.sub(r'[\s_()\-·/*,.×]', '', s).lower()


def numbers(s):
    return set(re.findall(r'\d+', s.replace(',', '')))


def score(a, b):
    na, nb = numbers(a), numbers(b)
    if na and nb and na != nb:          # 규격(숫자)이 다르면 다른 제품
        return 0.0
    return difflib.SequenceMatcher(None, norm(a), norm(b)).ratio()


def main(cat_path, master_path, ledger_path, out_path, unmatched_path):
    cat = list(csv.DictReader(open(cat_path, encoding='utf-8')))
    wb = openpyxl.load_workbook(master_path, data_only=True)
    master = []
    for sheet, label in (('03_단품마스터', '03 단품마스터'), ('04_선물세트마스터', '04 선물세트마스터')):
        ws = wb[sheet]
        for r in range(6, ws.max_row + 1):
            v = ws[f'C{r}'].value
            if v:
                master.append((label, str(v).strip()))
    sales = [r for r in csv.DictReader(open(ledger_path, encoding='utf-8-sig')) if r['kind'] == 'sales']
    cnt = collections.Counter(r['item_text'].strip() for r in sales)

    targets = []
    for c in cat:
        names = [c['카탈로그_제품명']]
        alias = re.sub(r'\s*\([A-Z]{2,3}-[^)]*\)\s*$', '', c['제품분류표_대응']).strip()
        if alias and not alias.startswith('없음'):
            names.append(alias)
        targets.append((c, names))

    rows = []
    for c, names in targets:
        base = dict(코드=c['코드'], 공식명=c['카탈로그_제품명'])
        rows.append({**base, '출처': '카탈로그', '다른표기': names[0], '장부건수': '', '판정': '공식명'})
        if len(names) > 1:
            rows.append({**base, '출처': '제품분류표', '다른표기': names[1], '장부건수': '', '판정': '확정'})
        for src, m in master:
            sc = max(score(m, x) for x in names)
            if sc >= 0.85:
                rows.append({**base, '출처': src, '다른표기': m, '장부건수': '', '판정': f'후보({sc:.2f})'})
    unmatched = []
    for name, c in cnt.most_common():
        sc, t = max(((max(score(name, x) for x in names), cc) for cc, names in targets), key=lambda z: z[0])
        if sc >= 0.8:
            rows.append(dict(코드=t['코드'], 공식명=t['카탈로그_제품명'], 출처='매출장부', 다른표기=name, 장부건수=c, 판정=f'후보({sc:.2f})'))
        else:
            unmatched.append((name, c, t['카탈로그_제품명'], round(sc, 2)))
    rows.sort(key=lambda r: (r['코드'], SRC_ORDER.index(r['출처']), -(r['장부건수'] or 0)))
    with open(out_path, 'w', encoding='utf-8-sig', newline='') as fp:
        w = csv.DictWriter(fp, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    with open(unmatched_path, 'w', encoding='utf-8-sig', newline='') as fp:
        w = csv.writer(fp)
        w.writerow(['장부_품목표기', '건수', '가장 비슷한 공식명', '유사도', '확정_코드(작성)'])
        w.writerows([u + ('',) for u in unmatched])
    tot = sum(cnt.values())
    covered = sum(r['장부건수'] for r in rows if r['출처'] == '매출장부')
    print(f'{len(rows)}줄 · 장부 표기 {len(cnt)}종 중 {len(cnt) - len(unmatched)}종 후보 연결 · 장부 줄 {covered / tot * 100:.1f}% · 미대응 {len(unmatched)}종')


if __name__ == '__main__':
    main(*sys.argv[1:6])
